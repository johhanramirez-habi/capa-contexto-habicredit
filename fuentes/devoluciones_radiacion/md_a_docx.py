#!/usr/bin/env python3
"""
md_a_docx.py — convierte un markdown de este repo a .docx

Fecha:  2026-08-13
Autor:  Claude (sesión Johhan Ramirez)

Cubre el subconjunto de markdown que usan los documentos de este proyecto:
encabezados, párrafos, tablas, listas con viñeta y numeradas, reglas horizontales,
**negrita**, *cursiva* y `código`. No pretende ser un conversor general.

Uso:  python3 md_a_docx.py RESUMEN_EJECUTIVO.md [salida.docx]
"""

import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

GRIS = RGBColor(0x44, 0x44, 0x44)
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


def escribir_inline(parrafo, texto):
    """Divide el texto en runs según **negrita**, `código` y *cursiva*."""
    for parte in INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            parrafo.add_run(parte[2:-2]).bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            run = parrafo.add_run(parte[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = GRIS
        elif parte.startswith("*") and parte.endswith("*"):
            parrafo.add_run(parte[1:-1]).italic = True
        else:
            parrafo.add_run(parte)


def abre_bloque(linea):
    """¿Esta línea empieza un bloque nuevo (encabezado, tabla, lista, regla)?"""
    d = linea.strip()
    return (
        not d
        or d.startswith(("#", "|"))
        or d in ("---", "***", "___")
        or bool(re.match(r"^([-*]\s+|\d+\.\s+)", d))
    )


def consumir(lineas, i, primera):
    """Junta un ítem de lista con sus líneas de continuación en un solo texto.

    En el markdown fuente los ítems largos se parten para respetar el ancho de
    columna; sin esto cada continuación terminaba como un párrafo suelto fuera de
    la lista.
    """
    partes = [primera]
    i += 1
    while i < len(lineas) and not abre_bloque(lineas[i]):
        partes.append(lineas[i].strip())
        i += 1
    return " ".join(partes), i


def celdas(linea):
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def es_separador(linea):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", linea.strip()))


def agregar_tabla(doc, bloque):
    encabezado = celdas(bloque[0])
    cuerpo = [celdas(l) for l in bloque[2:]]
    tabla = doc.add_table(rows=1, cols=len(encabezado))
    tabla.style = "Table Grid"
    tabla.autofit = True

    for i, texto in enumerate(encabezado):
        celda = tabla.rows[0].cells[i]
        celda.text = ""
        p = celda.paragraphs[0]
        escribir_inline(p, texto)
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(9)

    for fila in cuerpo:
        celdas_fila = tabla.add_row().cells
        for i, texto in enumerate(fila[: len(encabezado)]):
            celdas_fila[i].text = ""
            p = celdas_fila[i].paragraphs[0]
            escribir_inline(p, texto)
            for run in p.runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def convertir(ruta_md, ruta_docx):
    with open(ruta_md, encoding="utf-8") as f:
        lineas = f.read().split("\n")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for seccion in doc.sections:
        seccion.left_margin = seccion.right_margin = Inches(1.0)

    i = 0
    while i < len(lineas):
        linea = lineas[i]
        desnudo = linea.strip()

        if not desnudo:
            i += 1
            continue

        # Tabla: cabecera + separador + filas
        if desnudo.startswith("|") and i + 1 < len(lineas) and es_separador(lineas[i + 1]):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i])
                i += 1
            agregar_tabla(doc, bloque)
            continue

        if desnudo in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("• • •")
            run.font.color.rgb = GRIS
            i += 1
            continue

        if desnudo.startswith("#"):
            nivel = len(desnudo) - len(desnudo.lstrip("#"))
            texto = desnudo[nivel:].strip()
            encabezado = doc.add_heading(level=min(nivel, 4))
            escribir_inline(encabezado, texto)
            i += 1
            continue

        if re.match(r"^[-*]\s+", desnudo):
            texto, i = consumir(lineas, i, re.sub(r"^[-*]\s+", "", desnudo))
            escribir_inline(doc.add_paragraph(style="List Bullet"), texto)
            continue

        if re.match(r"^\d+\.\s+", desnudo):
            texto, i = consumir(lineas, i, re.sub(r"^\d+\.\s+", "", desnudo))
            escribir_inline(doc.add_paragraph(style="List Number"), texto)
            continue

        # Párrafo: junta líneas consecutivas hasta la próxima en blanco o marca
        buffer = []
        while i < len(lineas):
            actual = lineas[i].strip()
            if not actual or actual.startswith(("#", "|", "- ", "* ")) or actual in ("---",):
                break
            if re.match(r"^\d+\.\s+", actual):
                break
            buffer.append(actual)
            i += 1
        p = doc.add_paragraph()
        escribir_inline(p, " ".join(buffer))

    doc.save(ruta_docx)
    return ruta_docx


if __name__ == "__main__":
    entrada = sys.argv[1] if len(sys.argv) > 1 else "RESUMEN_EJECUTIVO.md"
    salida = sys.argv[2] if len(sys.argv) > 2 else entrada.rsplit(".", 1)[0] + ".docx"
    print("escrito ->", convertir(entrada, salida))
